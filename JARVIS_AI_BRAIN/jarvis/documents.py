"""Universal File Intelligence — hujjatlarni o'qish, tahlil, tahrirlash (NUSXADA), yaratish, solishtirish, OCR.

* O'qish: DOCX, PPTX, ODT/ODP/ODS (ZIP ichidagi XML — kutubxonasiz ham), RTF, PDF, XLSX, TXT/MD/CSV/JSON/kod, rasm (OCR).
* OCR: Tesseract o'rnatilgan bo'lsa (o'zbek/rus/ingliz). Bo'lmasa — aniq "OCR o'rnatilmagan" deyiladi, o'qigandek ko'rsatilmaydi.
* Tahrirlash: asl fayl HECH QACHON o'zgarmaydi — natija `exports/` da yangi fayl. Yaratish: DOCX, PPTX (python-docx/-pptx).
* Xavfsizlik: faqat ruxsat berilgan papkalar; ZIP ichidagi umumiy hajm cheklangan (zip-bomba); maxfiy fayllar o'qilmaydi;
  hujjat matni AI'ga faqat so'ralganda va MAXFIY rejimda (faqat mahalliy model) yuboriladi.
* Katta fayllar: bo'laklarga (chunk) bo'linadi; qidiruv kerakli bo'laklarni qaytaradi.
"""
from __future__ import annotations

import difflib
import re
import shutil
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from .security import contains_secret

MAX_FILE = 25 * 2**20
MAX_ZIP_TOTAL = 200 * 2**20
CHUNK = 4000
TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml", ".toml", ".xml", ".html", ".css", ".py", ".js", ".ts", ".tsx", ".jsx",
            ".java", ".c", ".cpp", ".php", ".sql", ".sh", ".ps1", ".bat", ".go", ".rs", ".ini", ".log"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
SUPPORTED = TEXT_EXT | IMAGE_EXT | {".docx", ".pptx", ".odt", ".odp", ".ods", ".rtf", ".pdf", ".xlsx", ".xlsm"}
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0"}


class DocError(Exception):
    pass


def chunks(text: str, size: int = CHUNK, overlap: int = 200) -> list[str]:
    if len(text) <= size:
        return [text] if text else []
    out, i = [], 0
    while i < len(text):
        out.append(text[i:i + size])
        i += size - overlap
    return out


def _slug(s: str) -> str:
    return re.sub(r"[^\w-]+", "_", s)[:40] or "fayl"


def _safe_zip(p: Path) -> zipfile.ZipFile:
    try:
        z = zipfile.ZipFile(p)
    except zipfile.BadZipFile:
        raise DocError("Fayl buzilgan yoki bu format emas (ZIP o'qilmadi)") from None
    if sum(i.file_size for i in z.infolist()) > MAX_ZIP_TOTAL:
        z.close()
        raise DocError("Hujjat ichidagi ma'lumot juda katta (zip-bomba himoyasi)")
    return z


def _xml(z: zipfile.ZipFile, name: str):
    try:
        return ET.fromstring(z.read(name))
    except KeyError:
        raise DocError(f"Hujjat tuzilishi buzilgan: {name} yo'q") from None
    except ET.ParseError:
        raise DocError("Hujjat XML'i buzilgan") from None


class Documents:
    def __init__(self, core):
        self.core = core

    @property
    def files(self):
        return self.core.files

    def _in(self, path: str) -> Path:
        p = self.files._check_read(path)
        if not p.is_file():
            raise DocError("Fayl topilmadi")
        if p.stat().st_size > MAX_FILE:
            raise DocError("Fayl 25 MB dan katta")
        return p

    def _out(self, src: Path, suffix: str, ext: str | None = None) -> Path:
        d = self.core.studio.exports
        return d / f"{src.stem}_{suffix}_{datetime.now():%Y%m%d_%H%M%S}{ext or src.suffix}"

    # ---------------- o'qish ----------------
    @staticmethod
    def _tesseract() -> str | None:
        """PATH'da yoki Windows'dagi odatiy o'rnatish joyida (PATH o'zgartirilmaydi)."""
        import os
        exe = shutil.which("tesseract")
        for base in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)"), os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs")):
            if not exe and base:
                cand = os.path.join(base, "Tesseract-OCR", "tesseract.exe")
                exe = cand if os.path.exists(cand) else None
        return exe

    def ocr_status(self) -> dict:
        exe = self._tesseract()
        langs = []
        if exe:
            try:
                langs = subprocess.run([exe, "--list-langs"], capture_output=True, text=True, timeout=10).stdout.split()[1:]
            except (OSError, subprocess.SubprocessError):
                langs = []
        return {"available": bool(exe), "languages": langs,
                "note": "" if exe else "OCR uchun Tesseract o'rnatilmagan (Windows: winget install UB-Mannheim.TesseractOCR). O'rnatilmaguncha rasm/skan matni o'qilmaydi."}

    def _ocr(self, p: Path) -> str:
        st = self.ocr_status()
        if not st["available"]:
            raise DocError(st["note"])
        lang = "+".join(x for x in ("uzb", "uzb_cyrl", "rus", "eng") if x in st["languages"]) or "eng"
        try:
            r = subprocess.run([self._tesseract(), str(p), "stdout", "-l", lang], capture_output=True, text=True, timeout=120)
        except subprocess.TimeoutExpired:
            raise DocError("OCR 120 soniyada tugamadi") from None
        if r.returncode != 0:
            raise DocError(f"OCR xatosi: {r.stderr.strip()[:200]}")
        return r.stdout

    def extract(self, path: str) -> dict:
        p = self._in(path)
        ext = p.suffix.lower()
        meta, tables, parts = {}, [], []
        if ext not in SUPPORTED:
            raise DocError(f"Qo'llab-quvvatlanmaydigan format: {ext or '(kengaytmasiz)'} — o'qilmadi")
        if ext == ".docx":
            with _safe_zip(p) as z:
                root = _xml(z, "word/document.xml")
            body = root.find("w:body", NS)
            for el in list(body) if body is not None else []:
                tag = el.tag.split("}")[1]
                if tag == "p":
                    t = "".join(x.text or "" for x in el.iter(f"{{{NS['w']}}}t"))
                    if t.strip():
                        parts.append(t)
                elif tag == "tbl":
                    rows = [["".join(x.text or "" for x in c.iter(f"{{{NS['w']}}}t")) for c in r.iter(f"{{{NS['w']}}}tc")] for r in el.iter(f"{{{NS['w']}}}tr")]
                    tables.append(rows)
                    parts.extend(" | ".join(r) for r in rows)
            meta = {"paragraphs": len(parts), "tables": len(tables)}
        elif ext == ".pptx":
            with _safe_zip(p) as z:
                names = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)), key=lambda n: int(re.findall(r"\d+", n)[0]))
                slides = []
                for i, n in enumerate(names, 1):
                    t = [x.text for x in _xml(z, n).iter(f"{{{NS['a']}}}t") if x.text and x.text.strip()]
                    slides.append({"n": i, "title": t[0] if t else "", "text": t[1:]})
                    parts.append(f"## Slayd {i}: " + (t[0] if t else "") + ("\n" + "\n".join(f"- {x}" for x in t[1:]) if len(t) > 1 else ""))
            meta = {"slides": len(slides), "outline": slides[:200]}
        elif ext in (".odt", ".odp", ".ods"):
            with _safe_zip(p) as z:
                root = _xml(z, "content.xml")
            for el in root.iter():
                if el.tag in (f"{{{NS['text']}}}p", f"{{{NS['text']}}}h"):
                    t = "".join(el.itertext()).strip()
                    if t:
                        parts.append(t)
            meta = {"paragraphs": len(parts)}
        elif ext == ".rtf":
            raw = p.read_text(encoding="latin-1", errors="replace")
            raw = re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1251", "replace"), raw)
            txt = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", raw)
            parts = [x.strip() for x in txt.splitlines() if x.strip()]
            meta = {"note": "RTF formatlashsiz o'qildi"}
        elif ext == ".pdf":
            txt = self.files._read_pdf(p)
            parts = [txt] if txt.strip() else []
            meta = {"note": "" if txt.strip() else "PDF'da matn qatlami yo'q (skan bo'lishi mumkin) — OCR uchun sahifalarni rasm sifatida yuboring"}
        elif ext in (".xlsx", ".xlsm"):
            parts = [self.files._read_xlsx(p)]
        elif ext in IMAGE_EXT:
            parts = [self._ocr(p)]
            meta = {"ocr": True}
        else:
            parts = [p.read_text(encoding="utf-8", errors="replace")]
        text = "\n".join(parts)
        ch = chunks(text)
        return {"path": str(p), "name": p.name, "type": ext.lstrip("."), "size": p.stat().st_size, "chars": len(text), "words": len(text.split()),
                "chunks": len(ch), "text": text[:CHUNK * 3], "truncated": len(text) > CHUNK * 3, "tables": tables[:20], "meta": meta,
                "has_secret": contains_secret(text),
                "actions": ["xulosa", "qidirish", "solishtirish", "matnga eksport"] + (["so'z almashtirish (nusxada)"] if ext == ".docx" else [])}

    def find(self, path: str, query: str, limit: int = 5) -> list[dict]:
        p = self._in(path)
        terms = [t.lower() for t in query.split() if len(t) > 1]
        scored = []
        for i, c in enumerate(chunks(self._full_text(p))):
            low = c.lower()
            s = sum(low.count(t) for t in terms)
            if s:
                at = min(low.find(t) for t in terms if t in low)       # parcha — topilgan so'z atrofidan
                scored.append({"chunk": i, "score": s, "text": c[max(0, at - 500):at + 1000]})
        return sorted(scored, key=lambda x: -x["score"])[:limit]

    def _full_text(self, p: Path) -> str:
        d = self.extract(str(p))
        if not d["truncated"]:
            return d["text"]
        # To'liq matn uchun qayta o'qish (extract ko'rsatish uchun qisqartiradi)
        ext = p.suffix.lower()
        if ext in TEXT_EXT:
            return p.read_text(encoding="utf-8", errors="replace")
        if ext == ".pdf":
            return self.files._read_pdf(p)
        return d["text"]

    def summary(self, path: str, use_ai: bool = False) -> dict:
        d = self.extract(path)
        sents = re.split(r"(?<=[.!?])\s+", d["text"].replace("\n", " "))
        extractive = " ".join(s.strip() for s in sents[:5] if s.strip())[:1500]
        out = {"name": d["name"], "words": d["words"], "chunks": d["chunks"], "extractive": extractive, "ai": None}
        if use_ai:
            prompt = f"Quyidagi hujjatning qisqa o'zbekcha xulosasini yoz (5 band):\n<untrusted>{d['text'][:6000]}</untrusted>"
            try:   # hujjat shaxsiy bo'lishi mumkin — faqat mahalliy model (private=True)
                r = self.core.ai.route([{"role": "user", "content": prompt}], purpose="document", private=True)
                out["ai"] = {"text": r["text"], "model": f"{r['route']['provider']}/{r['route']['model']}"}
            except Exception as e:
                out["ai"] = {"error": f"Mahalliy AI xulosa qila olmadi: {e}"}
        return out

    def compare(self, a: str, b: str) -> dict:
        ta, tb = self._full_text(self._in(a)), self._full_text(self._in(b))
        la, lb = ta.splitlines(), tb.splitlines()
        diff = list(difflib.unified_diff(la, lb, Path(a).name, Path(b).name, lineterm="", n=1))
        return {"similarity": round(difflib.SequenceMatcher(None, ta[:50000], tb[:50000]).ratio(), 3), "added": sum(1 for x in diff if x.startswith("+") and not x.startswith("+++")),
                "removed": sum(1 for x in diff if x.startswith("-") and not x.startswith("---")), "diff": diff[:400]}

    def export_text(self, path: str, fmt: str = "md") -> dict:
        if fmt not in ("md", "txt"):
            raise DocError("Format: md yoki txt")
        src = self._in(path)
        out = self._out(src, "matn", f".{fmt}")
        out.write_text(self._full_text(src), encoding="utf-8")
        return {"source": str(src), "output": str(out)}

    # ---------------- tahrirlash va yaratish (faqat yangi fayl) ----------------
    def docx_replace(self, path: str, replacements: dict[str, str]) -> dict:
        try:
            import docx
        except ImportError:
            raise DocError("Word tahrirlash uchun 'python-docx' o'rnatilmagan") from None
        src = self._in(path)
        if src.suffix.lower() != ".docx":
            raise DocError("Faqat .docx tahrirlanadi")
        if not replacements or any(not k for k in replacements):
            raise DocError("Almashtirishlar kerak: {eski: yangi}")
        if contains_secret(" ".join(replacements.values())):
            raise DocError("Yangi matnda maxfiy ma'lumot bor")
        doc = docx.Document(str(src))
        counts, rebuilt = dict.fromkeys(replacements, 0), 0

        def fix(par):
            nonlocal rebuilt
            for old, new in replacements.items():
                if old not in par.text:
                    continue
                n = par.text.count(old)
                done = False
                for run in par.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                        done = True
                if not done or old in par.text:          # ibora bir nechta "run"ga bo'lingan — paragraf qayta yig'iladi
                    full = par.text.replace(old, new)
                    for i, run in enumerate(par.runs):
                        run.text = full if i == 0 else ""
                    rebuilt += 1
                counts[old] += n
        for par in doc.paragraphs:
            fix(par)
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for par in cell.paragraphs:
                        fix(par)
        out = self._out(src, "tahrir")
        doc.save(str(out))
        self.core.audit.log("documents", f"Word nusxasi tahrirlandi: {src.name} → {out.name}", target=str(out), risk="B")
        return {"source": str(src), "output": str(out), "replaced": counts, "original_unchanged": True,
                "note": f"{rebuilt} paragrafda ibora bir necha formatlash bo'lagiga bo'lingan edi — u yerda formatlash soddalashdi" if rebuilt else ""}

    def create_docx(self, title: str, text: str) -> dict:
        try:
            import docx
        except ImportError:
            raise DocError("'python-docx' o'rnatilmagan") from None
        if contains_secret(text):
            raise DocError("Matnda maxfiy ma'lumot bor")
        doc = docx.Document()
        doc.add_heading(title or "Hujjat", 0)
        for block in (text or "").split("\n"):
            if block.startswith("# "):
                doc.add_heading(block[2:], 1)
            elif block.startswith(("- ", "* ")):
                doc.add_paragraph(block[2:], style="List Bullet")
            elif block.strip():
                doc.add_paragraph(block)
        out = self.core.studio.exports / f"{_slug(title or 'hujjat')}_{datetime.now():%Y%m%d_%H%M%S}.docx"
        doc.save(str(out))
        return {"output": str(out)}

    def create_pptx(self, title: str, slides: list[dict]) -> dict:
        try:
            from pptx import Presentation
        except ImportError:
            raise DocError("'python-pptx' o'rnatilmagan") from None
        if not slides or len(slides) > 60:
            raise DocError("1–60 ta slayd kerak")
        prs = Presentation()
        s0 = prs.slides.add_slide(prs.slide_layouts[0])
        s0.shapes.title.text = title or "Taqdimot"
        for sl in slides:
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = str(sl.get("title", ""))[:200]
            body = s.placeholders[1].text_frame
            bullets = [str(b)[:500] for b in (sl.get("bullets") or [])][:12]
            body.text = bullets[0] if bullets else ""
            for b in bullets[1:]:
                body.add_paragraph().text = b
        out = self.core.studio.exports / f"{_slug(title or 'taqdimot')}_{datetime.now():%Y%m%d_%H%M%S}.pptx"
        prs.save(str(out))
        return {"output": str(out), "slides": len(slides) + 1}

    def formats(self) -> dict:
        try:
            import docx  # noqa: F401
            w = True
        except ImportError:
            w = False
        try:
            import pptx  # noqa: F401
            pp = True
        except ImportError:
            pp = False
        return {"read": sorted(SUPPORTED), "ocr": self.ocr_status(), "write": {"docx_edit_copy": w, "docx_create": w, "pptx_create": pp},
                "not_supported": [".doc (eski Word)", ".ppt (eski PowerPoint)", ".xls (eski Excel)", "audio/video matni (Voice → Whisper orqali)"],
                "note": "Eski .doc/.ppt/.xls formatlarni Word/PowerPoint/Excel'da .docx/.pptx/.xlsx qilib saqlang."}
