"""v1.3: Agent Swarm + Supervisor, semantik xotira (soxta Ollama), File Intelligence (haqiqiy DOCX/PPTX/ODT/RTF), Computer Use (soxta drayver)."""
import json
import re
import zipfile

import pytest

from jarvis import net
from jarvis.computer_use import ComputerUseError
from jarvis.documents import DocError
from jarvis.swarm import SwarmError


# ---------------- Agent Swarm ----------------
def test_swarm_delegate_uses_receiver_permissions_and_supervisor(core):
    core.memory.add("Ombor shanba kuni yopiq", "knowledge", source="user")
    r = core.swarm.delegate("manager", "memory_agent", "memory.search", {"query": "ombor"})
    assert r["supervisor"]["ok"] and r["result"][0]["content"].startswith("Ombor")
    empty = core.swarm.delegate("manager", "memory_agent", "memory.search", {"query": "zzqqxx"})
    assert not empty["supervisor"]["ok"] and "bo'sh" in empty["supervisor"]["note"]
    with pytest.raises(SwarmError, match="chaqira olmaydi"):                 # qabul qiluvchi bu tool'ga ruxsatsiz
        core.swarm.delegate("manager", "instagram", "files.read", {"path": "x"})
    with pytest.raises(SwarmError, match="topshira olmaydi"):                # yuboruvchi o'zi qila olmaydigan ishni topshira olmaydi
        core.swarm.delegate("file_analysis", "computer", "computer.propose_command", {"command": "mkdir x"})
    ok = core.swarm.delegate("user", "computer", "computer.propose_command", {"command": "mkdir x"})
    assert ok["status"] == "approval_requested" and ok["supervisor"]["ok"]               # tashqi ta'sirli — faqat tasdiq so'rovi


def test_swarm_cycle_depth_and_chain_limits(core):
    core.memory.add("Test yozuvi", "knowledge", source="user")
    r = core.swarm.delegate("manager", "memory_agent", "memory.search", {"query": "test"}, chain="c1")
    assert r["chain"] == "c1"
    with pytest.raises(SwarmError, match="sikl"):
        core.swarm.delegate("manager", "memory_agent", "memory.search", {"query": "test"}, chain="c1")
    with pytest.raises(SwarmError, match="chuqur"):
        core.swarm.delegate("manager", "memory_agent", "memory.search", {"query": "x"}, chain="c2", depth=4)
    statuses = {h["status"] for h in core.swarm.history()}
    assert {"ok", "denied"} <= statuses


def test_mission_with_swarm_steps(core):
    core.memory.add("POS hisoboti har kuni 9:00 da tayyor", "knowledge", source="user")
    m = core.missions.create("xotiradan qidir: hisoboti keyin matnni tekshir: salom dunyo", use_ai=False)
    assert [n["agent"] for n in m["graph"]["nodes"]] == ["swarm", "swarm"]
    m = core.missions.start(m["id"], background=False)
    assert m["status"] == "completed", m["report"]
    assert all(s["verified"] == 1 for s in m["report"]["steps"])
    assert len(core.swarm.history()) == 2
    bad = core.missions.start(core.missions.create("xotiradan qidir: mutlaqoyoqsoz", use_ai=False)["id"], background=False)
    assert bad["status"] == "failed" and "Nazoratchi" in (bad["report"]["steps"][0]["error"] or "")


# ---------------- Semantik xotira ----------------
VOCAB = ["ombor", "sklad", "zaxira", "mahsulot", "savdo", "pul", "kassa", "dam", "yopiq", "shanba"]
SYN = {"sklad": "ombor", "zaxira": "ombor", "kassa": "pul"}


def _vec(text):
    words = re.findall(r"\w+", text.lower())
    v = [0.0] * len(VOCAB)
    for w in words:
        w = SYN.get(w, w)
        for i, k in enumerate(VOCAB):
            if w.startswith(k):
                v[i] += 1
    return v if any(v) else [0.01] * len(VOCAB)


class FakeOllama:
    def __init__(self, models=("nomic-embed-text:latest",)):
        self.models, self.pulled = list(models), []

    def __call__(self, method, url, data, headers, timeout):
        body = json.loads(data) if data else {}
        if url.endswith("/api/tags"):
            return 200, json.dumps({"models": [{"name": m} for m in self.models]})
        if url.endswith("/api/embeddings"):
            return 200, json.dumps({"embedding": _vec(body["prompt"])})
        if url.endswith("/api/pull"):
            self.pulled.append(body["model"])
            self.models.append(body["model"])
            return 200, json.dumps({"status": "success"})
        return 404, "{}"


def test_semantic_requires_explicit_consent_and_finds_by_meaning(core, monkeypatch):
    fake = FakeOllama(models=())
    monkeypatch.setattr(net, "TRANSPORT", fake)
    core.settings.save = lambda: None
    st = core.semantic.status()
    assert not st["enabled"] and not st["installed"] and "kalit so'z" in st["mode"]
    plan = core.semantic.install(confirm=False)
    assert plan["needs_confirm"] and plan["size_mb"] > 0 and not fake.pulled            # tasdiqsiz yuklanmaydi
    core.memory.add("Ombor shanba kuni yopiq bo'ladi", "knowledge", source="user")
    core.memory.add("Kassa hisoboti kechqurun", "knowledge", source="user")
    assert [m["match"] for m in core.semantic.hybrid("sklad qachon dam")] == []          # model yo'q — FTS so'z topmadi
    core.semantic.install(confirm=True)
    for _ in range(100):
        if not core.semantic.state["pulling"]:
            break
        __import__("time").sleep(0.05)
    assert fake.pulled == ["nomic-embed-text"] and core.settings.embedding_model == "nomic-embed-text"
    assert core.semantic.status()["indexed"] == 2
    hits = core.semantic.hybrid("sklad qachon dam", 3)
    assert hits and hits[0]["content"].startswith("Ombor") and hits[0]["match"] == "ma'no"
    mid = core.memory.add("Ombor shanba yopiq", "knowledge", source="user")["id"]
    core.semantic.index()
    assert core.semantic.similar(mid)                                                     # ma'no bo'yicha takror
    core.memory.delete(mid)
    core.semantic.index()
    assert core.db.one("SELECT COUNT(*) n FROM memory_vectors WHERE memory_id=?", (mid,))["n"] == 0


def test_semantic_ollama_down_falls_back(core, monkeypatch):
    def down(*a, **k):
        raise net.NetError("connection refused")
    monkeypatch.setattr(net, "TRANSPORT", down)
    core.settings.embedding_model = "nomic-embed-text"
    core.memory.add("Ombor kaliti", "knowledge", source="user")
    assert core.semantic.hybrid("ombor")[0]["match"] == "kalit so'z"
    assert core.semantic.index()["errors"] >= 1


# ---------------- File Intelligence ----------------
def _ws(tmp_path):
    return tmp_path / "workspace"


def test_docx_extract_edit_copy_create_and_compare(core, tmp_path):
    import docx
    ws = _ws(tmp_path)
    d = docx.Document()
    d.add_paragraph("Shartnoma raqami 15. Mijoz: Ali Valiyev.")
    p = d.add_paragraph()
    p.add_run("Narx: 100 ")
    p.add_run("000 so'm")                                     # ibora ikki "run"ga bo'lingan
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Mahsulot", "Soni", "Non", "5"
    src = ws / "shartnoma.docx"
    d.save(src)
    before = src.read_bytes()
    x = core.documents.extract(str(src))
    assert "Ali Valiyev" in x["text"] and x["tables"][0][1] == ["Non", "5"] and x["meta"]["tables"] == 1
    r = core.documents.docx_replace(str(src), {"Ali Valiyev": "Vali Aliyev", "100 000": "120 000"})
    assert src.read_bytes() == before and r["original_unchanged"]                       # asl fayl o'zgarmadi
    y = core.documents.extract(r["output"])
    assert "Vali Aliyev" in y["text"] and "120 000" in y["text"] and r["replaced"]["100 000"] == 1 and r["note"]
    cmp = core.documents.compare(str(src), r["output"])
    assert 0 < cmp["similarity"] < 1 and cmp["added"] >= 1
    new = core.documents.create_docx("Hisobot", "# Bo'lim\n- birinchi\nOddiy matn")
    assert "birinchi" in core.documents.extract(new["output"])["text"]
    ex = core.documents.export_text(str(src), "md")
    assert "Shartnoma" in open(ex["output"], encoding="utf-8").read()


def test_pptx_odt_rtf_and_honest_failures(core, tmp_path):
    ws = _ws(tmp_path)
    out = core.documents.create_pptx("Strategiya", [{"title": "Maqsad", "bullets": ["Savdoni oshirish", "Xarajatni kamaytirish"]}])
    ppt = core.documents.extract(out["output"])
    assert ppt["meta"]["slides"] == 2 and "Savdoni oshirish" in ppt["text"]
    odt = ws / "hujjat.odt"
    with zipfile.ZipFile(odt, "w") as z:
        z.writestr("content.xml", '<?xml version="1.0"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                   'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text><text:h>Sarlavha</text:h>'
                   '<text:p>ODT matni</text:p></office:text></office:body></office:document-content>')
    assert "ODT matni" in core.documents.extract(str(odt))["text"]
    rtf = ws / "x.rtf"
    rtf.write_text(r"{\rtf1\ansi{\fonttbl\f0 Arial;}\f0 Salom RTF\par}", encoding="latin-1")
    assert "Salom RTF" in core.documents.extract(str(rtf))["text"]
    bad = ws / "buzuq.docx"
    bad.write_bytes(b"bu docx emas")
    with pytest.raises(DocError, match="buzilgan"):
        core.documents.extract(str(bad))
    bomb = ws / "bomba.docx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * (210 * 2**20))
    with pytest.raises(DocError, match="zip-bomba"):
        core.documents.extract(str(bomb))
    (ws / "eski.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(DocError, match="Qo'llab-quvvatlanmaydigan"):
        core.documents.extract(str(ws / "eski.doc"))
    img = ws / "skan.png"
    img.write_bytes(b"\x89PNG\r\n")
    if not core.documents.ocr_status()["available"]:
        with pytest.raises(DocError, match="Tesseract"):                                # OCR yo'q — o'qigandek ko'rsatilmaydi
            core.documents.extract(str(img))
    with pytest.raises(Exception):
        core.documents.extract(str(tmp_path / "boshqa_joy.txt"))                        # ruxsat berilmagan papka
    assert ".doc (eski Word)" in core.documents.formats()["not_supported"][0]


def test_large_text_chunking_and_find(core, tmp_path):
    ws = _ws(tmp_path)
    big = ws / "katta.txt"
    big.write_text(("oddiy qator matni. " * 300 + "\n") * 10 + "MAXSUS kalit jumla shu yerda.\n", encoding="utf-8")
    x = core.documents.extract(str(big))
    assert x["chunks"] > 5 and x["truncated"]
    hits = core.documents.find(str(big), "maxsus kalit")
    assert hits and "MAXSUS" in hits[0]["text"]


# ---------------- Computer Use ----------------
class FakeDriver:
    def __init__(self):
        self.calls, self.title = [], "Bloknot — hisobot.txt"

    def size(self):
        return (1920, 1080)

    def screenshot(self, path):
        from PIL import Image
        Image.new("RGB", (10, 10)).save(path)

    def click(self, x, y, button="left", clicks=1):
        self.calls.append(("click", x, y, button, clicks))

    def move(self, x, y):
        self.calls.append(("move", x, y))

    def type_text(self, text):
        self.calls.append(("type", text))

    def hotkey(self, keys):
        self.calls.append(("hotkey", tuple(keys)))

    def scroll(self, n):
        self.calls.append(("scroll", n))

    def windows(self):
        return [{"title": self.title, "active": True}, {"title": "Payme — kirish", "active": False}]

    def active_title(self):
        return self.title

    def focus(self, title):
        self.calls.append(("focus", title))

    def ui_tree(self, title, limit=150):
        return [{"name": "Saqlash", "type": "Button", "x": 10, "y": 20}]


def _approve(core, appr):
    core.approvals.decide(appr["id"], True, appr["payload_hash"])


def test_computer_use_requires_approval_and_blocks_sensitive(core):
    drv = FakeDriver()
    core.computer_use.driver = drv
    assert core.computer_use.status()["available"]
    wins = core.computer_use.windows()
    assert any(w["sensitive"] for w in wins) and not all(w["sensitive"] for w in wins)
    acts = [{"type": "focus_window", "title": "Bloknot"}, {"type": "click", "x": 100, "y": 200}, {"type": "type", "text": "Salom"},
            {"type": "hotkey", "keys": ["ctrl", "s"]}]
    p = core.computer_use.propose(acts, "hisobotni saqlash")
    assert p["approval"]["status"] == "pending" and not drv.calls                       # tasdiqsiz hech narsa bajarilmadi
    with pytest.raises(Exception):
        core.computer_use.execute(p["approval"]["id"])
    _approve(core, p["approval"])
    r = core.computer_use.execute(p["approval"]["id"])
    assert r["done"] == 4 and r["verified"] and drv.calls[2] == ("type", "Salom")
    with pytest.raises(Exception):
        core.computer_use.execute(p["approval"]["id"])                                  # ikki marta bajarilmaydi
    assert len(core.computer_use.log()) == 4
    for bad in ([{"type": "type", "text": "token sk-" + "a" * 40}], [{"type": "hotkey", "keys": ["ctrl", "alt", "delete"]}],
                [{"type": "focus_window", "title": "Payme kirish"}], [{"type": "rm"}], [{"type": "click", "x": -5, "y": 1}]):
        with pytest.raises(ComputerUseError):
            core.computer_use.propose(bad)


def test_computer_use_stops_on_sensitive_window_and_error(core):
    drv = FakeDriver()
    core.computer_use.driver = drv
    p = core.computer_use.propose([{"type": "click", "x": 1, "y": 1}, {"type": "type", "text": "matn"}])
    _approve(core, p["approval"])
    drv.title = "Bank kirish — parol"                                                    # bajarish paytida maxfiy oyna ochildi
    r = core.computer_use.execute(p["approval"]["id"])
    assert r["done"] == 0 and "maxfiy" in r["error"] and not drv.calls
    assert core.approvals.get(p["approval"]["id"])["status"] == "unknown"
    with pytest.raises(ComputerUseError):
        core.computer_use.screenshot()


def test_computer_use_macros_screenshot_and_limits(core):
    drv = FakeDriver()
    core.computer_use.driver = drv
    core.computer_use.save_macro("saqlash", [{"type": "hotkey", "keys": ["ctrl", "s"]}])
    rp = core.computer_use.replay("saqlash")
    assert rp["approval"]["status"] == "pending"                                         # qayta ijro ham yangi tasdiq bilan
    shot = core.computer_use.screenshot()
    assert shot["path"].endswith(".png")
    core.settings.autonomy_level = 0
    with pytest.raises(ComputerUseError, match="Level 0"):
        core.computer_use.propose([{"type": "wait", "seconds": 1}])
    core.settings.autonomy_level = 2
    p = core.computer_use.propose([{"type": "wait", "seconds": 0.1}])
    _approve(core, p["approval"])
    core.automation_paused = True
    with pytest.raises(ComputerUseError, match="Favqulodda"):
        core.computer_use.execute(p["approval"]["id"])
    core.automation_paused = False
    core.computer_use.driver = None
    core.computer_use._driver_error = None
    if __import__("sys").platform != "win32":
        assert not core.computer_use.status()["available"] and "Windows" in core.computer_use.status()["note"]


def test_v13_api(client, core, tmp_path, monkeypatch):
    assert client.get("/api/swarm").status_code == 200
    assert client.post("/api/swarm/delegate", json={"from_agent": "manager", "to_agent": "x", "tool": "nope"}).status_code == 400
    monkeypatch.setattr(net, "TRANSPORT", FakeOllama(models=()))
    assert client.post("/api/semantic/install").json()["needs_confirm"]
    assert client.post("/api/semantic/install?model=bad;rm").status_code == 400
    assert ".docx" in client.get("/api/docs/formats").json()["read"]
    r = client.post("/api/docs-upload", files={"file": ("eslatma.txt", b"Salom JARVIS hujjati", "text/plain")})
    assert r.status_code == 200 and "Salom" in r.json()["text"]
    assert client.post("/api/docs-upload", files={"file": ("virus.exe", b"MZ", "application/octet-stream")}).status_code == 400
    assert client.post("/api/docs/explode", json={}).status_code == 404
    core.computer_use.driver = FakeDriver()
    assert client.get("/api/computer-use").json()["available"]
    pr = client.post("/api/computer-use/propose", json={"actions": [{"type": "wait", "seconds": 0.1}]}).json()
    assert pr["approval"]["status"] == "pending"
    assert client.post(f"/api/computer-use/execute/{pr['approval']['id']}").status_code == 400      # tasdiqlanmagan


def test_tool_registry_has_new_tools(core):
    ids = {t["id"] for t in core.tools.catalog()}
    assert {"docs.extract", "docs.find", "computer.screenshot", "computer.propose_actions"} <= ids
